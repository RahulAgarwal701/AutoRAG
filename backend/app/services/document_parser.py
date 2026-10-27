import re
from io import BytesIO
from pathlib import Path
from statistics import median

from app.services.ingestion_models import ParsedBlock, ParsedDocument

_MARKDOWN_HEADING = re.compile(r"^#{1,6}\s+(.+?)\s*$")


def _markdown_blocks(text: str) -> list[ParsedBlock]:
    blocks: list[ParsedBlock] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            value = "\n".join(buffer).strip()
            if value:
                blocks.append(ParsedBlock(text=value, section=section))
            buffer.clear()

    for line in text.splitlines():
        heading = _MARKDOWN_HEADING.match(line)
        if heading:
            flush()
            section = heading.group(1).strip()
        else:
            buffer.append(line)
    flush()
    return blocks


def _pdf_blocks(data: bytes) -> list[ParsedBlock]:
    import fitz

    blocks: list[ParsedBlock] = []
    try:
        with fitz.open(stream=data, filetype="pdf") as pdf:
            for page_number, page in enumerate(pdf, start=1):
                page_data = page.get_text("dict")
                text_blocks = [block for block in page_data.get("blocks", []) if "lines" in block]
                font_sizes = [
                    float(span.get("size", 0))
                    for block in text_blocks
                    for line in block["lines"]
                    for span in line["spans"]
                    if str(span.get("text", "")).strip()
                ]
                body_size = median(font_sizes) if font_sizes else 0
                section: str | None = None
                for block in text_blocks:
                    spans = [
                        span
                        for line in block["lines"]
                        for span in line["spans"]
                        if str(span.get("text", "")).strip()
                    ]
                    text = " ".join(str(span["text"]).strip() for span in spans).strip()
                    if not text:
                        continue
                    largest_size = max(float(span.get("size", 0)) for span in spans)
                    looks_like_heading = len(text) <= 160 and largest_size > body_size * 1.15
                    if looks_like_heading:
                        section = text
                    else:
                        blocks.append(ParsedBlock(text=text, page=page_number, section=section))
    except Exception:
        from pdfminer.high_level import extract_text

        text = extract_text(BytesIO(data)).strip()
        if text:
            blocks.append(ParsedBlock(text=text))
    return blocks


def _docx_blocks(data: bytes) -> list[ParsedBlock]:
    from docx import Document

    blocks: list[ParsedBlock] = []
    section: str | None = None
    for paragraph in Document(BytesIO(data)).paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            section = text
        else:
            blocks.append(ParsedBlock(text=text, section=section))
    return blocks


def _html_blocks(data: bytes) -> list[ParsedBlock]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    blocks: list[ParsedBlock] = []
    section: str | None = None
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li"]):
        text = element.get_text(" ", strip=True)
        if not text:
            continue
        if element.name and element.name.startswith("h"):
            section = text
        else:
            blocks.append(ParsedBlock(text=text, section=section))
    return blocks


def parse_document_bytes(data: bytes, filename: str) -> ParsedDocument:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        blocks = _pdf_blocks(data)
    elif extension == ".docx":
        blocks = _docx_blocks(data)
    elif extension in {".html", ".htm"}:
        blocks = _html_blocks(data)
    elif extension == ".md":
        blocks = _markdown_blocks(data.decode("utf-8", errors="replace"))
    elif extension == ".txt":
        text = data.decode("utf-8", errors="replace").strip()
        blocks = [ParsedBlock(text=text)] if text else []
    else:
        raise ValueError(f"Unsupported document extension: {extension}")
    if not blocks:
        raise ValueError("The document did not contain extractable text")
    return ParsedDocument(blocks=blocks)
